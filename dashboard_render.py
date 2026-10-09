"""Static renderings of the dashboard: PNG figures (matplotlib) and the PowerPoint report (python-pptx)."""
import io
import math
from collections import OrderedDict
from datetime import datetime

import numpy as np
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure

CARD_TITLES = OrderedDict([
    ("signals", "Tracked Signals - x(t), y(t)"),
    ("periods", "Periods and Frequencies"),
    ("amplitudes", "Amplitudes and Damping"),
    ("shape", "Orientation, Shape and Classification"),
    ("spectral", "Spectral Check and Independent Damping Fit"),
    ("sliding", "Sliding-Window Evolution"),
    ("phasespace", "Phase Space"),
    ("stats", "Statistical Summary"),
    ("histograms", "Histograms and Distributions"),
    ("correlation", "Correlation Matrix"),
    ("velocity", "Velocity and Acceleration"),
    ("energy", "Energy Proxy and Decay"),
    ("residuals", "Residuals and Outlier Detection"),
])
BLUE, GREEN, AMBER, RED, VIOLET, PINK = "#3b82f6", "#10b981", "#f59e0b", "#ef4444", "#a78bfa", "#f472b6"


# ----------------------------------------------------------------------------- figure helpers
class _Theme:
    def __init__(self, dark):
        self.fig = "#1a2332" if dark else "#ffffff"
        self.ax = "#0f172a" if dark else "#f8fafc"
        self.fg = "#e7ecf3" if dark else "#1e293b"
        self.grid = "#2d3a4f" if dark else "#d5dde8"


def _new(theme, rows=1, cols=1, size=(10, 5.2), **kw):
    fig = Figure(figsize=size, dpi=130, facecolor=theme.fig)
    FigureCanvasAgg(fig)
    axes = fig.subplots(rows, cols, **kw)
    return fig, axes


def _style(theme, *axes, title=None, xlabel=None, ylabel=None, grid=True):
    for ax in axes:
        ax.set_facecolor(theme.ax)
        ax.tick_params(colors=theme.fg, labelsize=8)
        for s in ax.spines.values():
            s.set_color(theme.grid)
        if grid:
            ax.grid(True, color=theme.grid, lw=0.6, alpha=0.8)
        ax.xaxis.label.set_color(theme.fg)
        ax.yaxis.label.set_color(theme.fg)
        ax.title.set_color(theme.fg)
        leg = ax.get_legend()
        if leg:
            leg.get_frame().set_facecolor(theme.ax)
            leg.get_frame().set_edgecolor(theme.grid)
            for t in leg.get_texts():
                t.set_color(theme.fg)
    a0 = axes[0]
    if title:
        a0.set_title(title, fontsize=10)
    if xlabel:
        a0.set_xlabel(xlabel, fontsize=9)
    if ylabel:
        a0.set_ylabel(ylabel, fontsize=9)


def _png(fig):
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", facecolor=fig.get_facecolor())
    return buf.getvalue()


def _thin(a, k):
    return a[::k] if k > 1 else a


def _k(n, cap=6000):
    return max(1, int(math.ceil(n / cap)))


def _g(v, d=3):
    return "-" if v is None or not np.isfinite(v) else f"{v:.{d}f}"


def _no_data(theme, text):
    fig, ax = _new(theme)
    ax.axis("off")
    ax.text(0.5, 0.5, text, ha="center", va="center", color=theme.fg, fontsize=12, transform=ax.transAxes)
    return _png(fig)


# ----------------------------------------------------------------------------- figures
def _fig_signals(R, th):
    F, u = R["full"], R["unit"]
    k = _k(R["n"])
    t = _thin(F["t"], k)
    fig, ax = _new(th)
    ax.plot(t, _thin(F["xs"], k), color=BLUE, lw=0.6, ls=":", alpha=0.55, label="x original")
    ax.plot(t, _thin(F["ys"], k), color=GREEN, lw=0.6, ls=":", alpha=0.55, label="y original")
    ax.plot(t, _thin(F["x"], k), color=BLUE, lw=1.1, label="x filtered")
    ax.plot(t, _thin(F["y"], k), color=GREEN, lw=1.1, label="y filtered")
    ax.legend(ncol=4, fontsize=8, loc="upper right")
    _style(th, ax, title=f"{R['series_label']} - {R['filter_desc']}", xlabel="Time (s)", ylabel=f"Position ({u})")
    return _png(fig)


def _fig_periods(R, th):
    S, P = R["spectrum"], R["periods"]
    fig, ax = _new(th)
    ax.plot(S["f"], S["x"], color=BLUE, lw=1, label="X spectrum")
    ax.plot(S["f"], S["y"], color=GREEN, lw=1, label="Y spectrum")
    for v, c in ((P["fx_fft"], BLUE), (P["fy_fft"], GREEN)):
        if np.isfinite(v):
            ax.axvline(v, color=c, ls="--", lw=0.8, alpha=0.8)
    ax.set_xlim(0, S["range_hi"])
    ax.legend(fontsize=8)
    _style(th, ax, title="Amplitude spectrum (Hann window, zero padded)", xlabel="Frequency (Hz)",
           ylabel=f"Amplitude ({R['unit']})")
    return _png(fig)


def _fig_amplitudes(R, th):
    A = R["amp"]
    if len(A["env_t"]) < 3:
        return _no_data(th, "Not enough data for an amplitude envelope")
    fig, ax = _new(th)
    ax.plot(A["env_t"], A["env_a"], "o", ms=3, color=AMBER, label="Envelope (sliding RMS)")
    if np.isfinite(A["gamma"]) and np.isfinite(A["A0"]):
        ax.plot(A["env_t"], A["A0"] * np.exp(-A["gamma"] * np.asarray(A["env_t"])), color=RED, lw=2,
                label=f"A0 exp(-gamma t), gamma={A['gamma']:.5f}/s")
    ax.legend(fontsize=8)
    _style(th, ax, title="Amplitude envelope and exponential damping", xlabel="Time (s)",
           ylabel=f"Amplitude ({R['unit']})")
    return _png(fig)


def _fig_shape(R, th):
    F, u = R["full"], R["unit"]
    k = _k(R["n"], 8000)
    x, y = _thin(F["x"], k), _thin(F["y"], k)
    fig, ax = _new(th, size=(7, 6))
    ax.plot(x, y, color=VIOLET, lw=0.8)
    ax.plot(x[0], y[0], "o", color=GREEN, label="start")
    ax.plot(x[-1], y[-1], "o", color=RED, label="end")
    ax.set_aspect("equal", adjustable="datalim")
    ax.legend(fontsize=8)
    _style(th, ax, title=f"Trajectory - {R['shape']['label']}", xlabel=f"x ({u})", ylabel=f"y ({u})")
    return _png(fig)


def _fig_spectral(R, th):
    S, D = R["spectrum"], R["damping"]
    fig, (a1, a2) = _new(th, 1, 2, size=(11, 4.8))
    a1.semilogy(S["f"], np.maximum(S["x"], 1e-9), color=BLUE, lw=1, label="X")
    a1.semilogy(S["f"], np.maximum(S["y"], 1e-9), color=GREEN, lw=1, label="Y")
    a1.set_xlim(0, S["range_hi"])
    a1.legend(fontsize=8)
    _style(th, a1, title="Spectrum (log scale)", xlabel="f (Hz)", ylabel="|FFT| amplitude")
    if len(D["peaks_t"]) >= 3:
        a2.plot(D["peaks_t"], D["peaks_a"], "o", ms=3, color=AMBER, label="peaks of r(t)")
        row = next((r for r in D["rows"] if r["method"].startswith("Peak")), None)
        if row and np.isfinite(row["gamma"]):
            tt = np.asarray(D["peaks_t"])
            a2.plot(tt, row["A0"] * np.exp(-row["gamma"] * tt), color=RED, lw=2,
                    label=f"fit gamma={row['gamma']:.5f}/s")
        a2.legend(fontsize=8)
    _style(th, a2, title="Independent damping fit", xlabel="Time (s)", ylabel=f"r ({R['unit']})")
    return _png(fig)


def _fig_sliding(R, th):
    S = R["sliding"]
    if len(S["t"]) < 2:
        return _no_data(th, f"Need a recording longer than {S['need']:.0f} s for sliding windows")
    fig, (a1, a2) = _new(th, 2, 1, size=(10, 6), sharex=True)
    a1.plot(S["t"], S["Tx"], color=BLUE, lw=1, label="Tx")
    a1.plot(S["t"], S["Ty"], color=GREEN, lw=1, label="Ty")
    a1.legend(fontsize=8)
    a2.plot(S["t"], S["theta"], color=VIOLET, lw=1)
    _style(th, a1, title=f"Sliding window {S['win']:.0f} s", ylabel="Period (s)")
    _style(th, a2, xlabel="Time (s)", ylabel="theta (deg)")
    return _png(fig)


def _fig_phase(R, th):
    F, u = R["full"], R["unit"]
    k = _k(R["n"], 5000)
    fig, ax = _new(th, size=(7, 6))
    ax.plot(_thin(F["x"], k), _thin(F["vx"], k), color=PINK, lw=0.7, label="x - vx")
    ax.plot(_thin(F["y"], k), _thin(F["vy"], k), color=BLUE, lw=0.7, alpha=0.8, label="y - vy")
    ax.legend(fontsize=8)
    _style(th, ax, title="Phase portrait", xlabel=f"position ({u})", ylabel=f"velocity ({u}/s)")
    return _png(fig)


def _fig_stats(R, th):
    u = R["unit"]
    rows = []
    for key, lab in (("x", f"x ({u})"), ("y", f"y ({u})"), ("vx", f"vx ({u}/s)"), ("vy", f"vy ({u}/s)"),
                     ("ax", f"ax ({u}/s2)"), ("ay", f"ay ({u}/s2)")):
        s = R["stats"][key]
        rows.append([lab] + [_g(s[m], 2) for m in ("mean", "std", "min", "max", "median", "rms")])
    fig, ax = _new(th, size=(10, 3.4))
    ax.axis("off")
    tb = ax.table(cellText=rows, colLabels=["Signal", "Mean", "SD", "Min", "Max", "Median", "RMS"], loc="center",
                  cellLoc="left")
    tb.auto_set_font_size(False)
    tb.set_fontsize(9)
    tb.scale(1, 1.6)
    for (r, c), cell in tb.get_celld().items():
        cell.set_facecolor(th.ax if r else th.grid)
        cell.set_edgecolor(th.grid)
        cell.get_text().set_color(th.fg)
    return _png(fig)


def _fig_hist(R, th):
    fig, ax = _new(th)
    for key, c in (("x", BLUE), ("y", GREEN)):
        h = R["hist"][key]
        ax.bar(h["c"], h["n"], width=h["w"], color=c, alpha=0.65, label=key)
    ax.legend(fontsize=8)
    _style(th, ax, title="Position distribution", xlabel=f"Position ({R['unit']})", ylabel="Count")
    return _png(fig)


def _fig_corr(R, th):
    labels, m = R["corr"]["labels"], np.asarray(R["corr"]["matrix"], float)
    fig, ax = _new(th, size=(6.2, 5.4))
    im = ax.imshow(m, vmin=-1, vmax=1, cmap="RdBu_r")
    ax.set_xticks(range(4), labels)
    ax.set_yticks(range(4), labels)
    for i in range(4):
        for j in range(4):
            ax.text(j, i, _g(m[i, j], 2), ha="center", va="center", fontsize=10,
                    color="black" if abs(m[i, j]) < 0.6 or not np.isfinite(m[i, j]) else "white")
    cb = fig.colorbar(im, ax=ax)
    cb.ax.tick_params(colors=th.fg, labelsize=8)
    _style(th, ax, title="Pearson correlation", grid=False)
    return _png(fig)


def _fig_velocity(R, th):
    F, u = R["full"], R["unit"]
    k = _k(R["n"])
    t = _thin(F["t"], k)
    fig, (a1, a2) = _new(th, 2, 1, size=(10, 6.4), sharex=True)
    a1.plot(t, _thin(F["vx"], k), color=BLUE, lw=0.8, label="vx")
    a1.plot(t, _thin(F["vy"], k), color=GREEN, lw=0.8, label="vy")
    a1.legend(fontsize=8)
    a2.plot(t, _thin(F["ax"], k), color=AMBER, lw=0.8, label="ax")
    a2.plot(t, _thin(F["ay"], k), color=RED, lw=0.8, alpha=0.8, label="ay")
    a2.legend(fontsize=8)
    _style(th, a1, ylabel=f"velocity ({u}/s)", title="Velocity and acceleration (central differences)")
    _style(th, a2, xlabel="Time (s)", ylabel=f"acceleration ({u}/s2)")
    return _png(fig)


def _fig_energy(R, th):
    F = R["full"]
    k = _k(R["n"])
    t = _thin(F["t"], k)
    fig, ax = _new(th)
    ax.plot(t, _thin(F["ke"], k), color=BLUE, lw=0.7, label="KE proxy")
    ax.plot(t, _thin(F["pe"], k), color=GREEN, lw=0.7, label="PE proxy")
    ax.plot(t, _thin(F["te"], k), color=AMBER, lw=1, label="Total")
    ax.plot(t, _thin(F["tes"], k), color=RED, lw=2, label="Total (1-period average)")
    ax.legend(fontsize=8, ncol=4)
    _style(th, ax, title="Energy proxy: 0.5 v^2 + 0.5 w0^2 r^2", xlabel="Time (s)", ylabel="Energy proxy (arb. units)")
    return _png(fig)


def _fig_resid(R, th):
    F = R["full"]
    k = _k(R["n"])
    t = _thin(F["t"], k)
    fig, ax = _new(th)
    ax.plot(t, _thin(F["rx"], k), color=RED, lw=0.7, label="x filtered - original")
    ax.plot(t, _thin(F["ry"], k), color=VIOLET, lw=0.7, label="y filtered - original")
    ax.legend(fontsize=8)
    _style(th, ax, title="Filter residuals", xlabel="Time (s)", ylabel=f"Residual ({R['unit']})")
    return _png(fig)


_RENDERERS = OrderedDict([
    ("signals", _fig_signals), ("periods", _fig_periods), ("amplitudes", _fig_amplitudes), ("shape", _fig_shape),
    ("spectral", _fig_spectral), ("sliding", _fig_sliding), ("phasespace", _fig_phase), ("stats", _fig_stats),
    ("histograms", _fig_hist), ("correlation", _fig_corr), ("velocity", _fig_velocity), ("energy", _fig_energy),
    ("residuals", _fig_resid),
])


def render_all(R, dark=True):
    th = _Theme(dark)
    out = OrderedDict()
    for cid, fn in _RENDERERS.items():
        try:
            out[cid] = fn(R, th)
        except Exception as e:  # noqa: BLE001 - one broken chart must not break the whole export
            out[cid] = _no_data(th, f"Could not render this chart: {e}")
    return out


# ----------------------------------------------------------------------------- PowerPoint
def _slide_content(R):
    """cid -> (bullets, table rows or None, notes)."""
    u, P, A, S, SH, D, E, RS = (R["unit"], R["periods"], R["amp"], R["sliding"], R["shape"], R["damping"],
                                R["energy"], R["resid"])
    cov = f"{R['coverage'] * 100:.1f}%" if R["coverage"] is not None else "n/a"
    c = OrderedDict()
    c["signals"] = ([
        f"Series: {R['series_label']}",
        f"Duration {R['duration']:.3f} s, {R['n']:,} samples, fs = {R['fs']:.3f} Hz, tracking coverage {cov}",
        f"Filter: {R['filter_desc']}",
        f"Outlier removal: {R['params']['outlier']} ({R['outliers']['x']} x / {R['outliers']['y']} y samples replaced)",
    ], None, "x(t) and y(t) after outlier removal and filtering, with the original series dotted.")
    c["periods"] = ([
        f"Peak method: Tx = {_g(P['Tx'])} s, Ty = {_g(P['Ty'])} s",
        f"Peak frequencies: fx = {_g(P['fx_peak'], 4)} Hz, fy = {_g(P['fy_peak'], 4)} Hz",
        f"FFT frequencies: fx = {_g(P['fx_fft'], 4)} Hz, fy = {_g(P['fy_fft'], 4)} Hz",
        "Peak and FFT estimates can differ (resolution, harmonics, non-stationarity) - report the method.",
    ], [["Metric", "Peak", "FFT"],
        ["fx (Hz)", _g(P["fx_peak"], 4), _g(P["fx_fft"], 4)], ["fy (Hz)", _g(P["fy_peak"], 4), _g(P["fy_fft"], 4)],
        ["df (Hz)", _g(P["df_peak"], 4), _g(P["df_fft"], 4)], ["fx/fy", _g(P["ratio_peak"], 4), _g(P["ratio_fft"], 4)]],
        "Spectrum is amplitude-scaled with a Hann window and zero padding.")
    c["amplitudes"] = ([
        f"Ax = {_g(A['Ax'], 2)} {u}, Ay = {_g(A['Ay'], 2)} {u}, Ax/Ay = {_g(A['ratio'], 2)}",
        f"Envelope A0 = {_g(A['A0'], 2)} {u}, gamma = {_g(A['gamma'], 5)} 1/s, tau = {_g(A['tau'], 1)} s",
        f"Quality factor Q = pi / (gamma T) = {_g(A['Q'], 1)}",
        f"Amplitude decay first to last third: {_g(A['decay'] * 100 if np.isfinite(A['decay']) else A['decay'], 1)} %",
    ], None, "Model A(t) = A0 exp(-gamma t) fitted to the sliding-RMS amplitude envelope.")
    c["shape"] = ([
        f"Classification: {SH['label']} - {SH['detail']}",
        f"theta mean {_g(SH['theta'], 2)} deg, drift {_g(SH['drift'], 2)} deg/s, minor/major {_g(SH['ecc'], 3)}",
        f"x-y phase {_g(SH['phase'], 1)} deg",
        f"Equivalent simple-pendulum length L = g (T/2pi)^2 = {_g(SH['L_eq'], 3)} m (g = {SH['g']:.2f})",
    ] + ([f"Theory T = {_g(SH['T_theory'])} s, relative error {_g(SH['T_err'] * 100, 1)} %"]
         if np.isfinite(SH["T_theory"]) else []), None, "Shape from amplitudes, phase, frequency offset and PCA.")
    c["spectral"] = ([
        f"fx FFT = {_g(P['fx_fft'], 4)} Hz, fy FFT = {_g(P['fy_fft'], 4)} Hz, ratio {_g(P['ratio_fft'], 4)}",
        "Damping estimated independently by several methods (table).",
    ], [["Method", "gamma (1/s)", "tau (s)", "R2"]] +
        [[r["method"], _g(r["gamma"], 5), _g(r["tau"], 1), _g(r["r2"], 3)] for r in D["rows"]],
        "Independent fits show whether the decay is a single exponential.")
    c["sliding"] = ([
        f"Window {S['win']:.0f} s, step {S['step']:.2f} s",
        f"Tx median {_g(S['Tx_med'])} s (SD {_g(S['Tx_sd'])}), Ty median {_g(S['Ty_med'])} s (SD {_g(S['Ty_sd'])})",
        f"theta drift {_g(S['drift'], 2)} deg/s, median minor/major {_g(S['ecc_med'], 3)}",
    ], None, "Period and orientation evolution over time.")
    c["phasespace"] = ([
        f"x SD {_g(R['stats']['x']['std'], 2)} {u}, y SD {_g(R['stats']['y']['std'], 2)} {u}",
        f"Peak speed {_g(R['speed']['peak'], 1)} {u}/s, median speed {_g(R['speed']['median'], 1)} {u}/s",
        "Damping makes the phase portrait spiral inward.",
    ], None, "Velocity from central differences of the filtered series.")
    st = R["stats"]
    c["stats"] = ([f"Computed over {R['n']:,} filtered samples."],
                  [["Signal", "Mean", "SD", "Min", "Max", "Median"]] +
                  [[k] + [_g(st[k][m], 2) for m in ("mean", "std", "min", "max", "median")] for k in st],
                  "Descriptive statistics for position, velocity and acceleration.")
    c["histograms"] = ([f"x range {_g(st['x']['min'], 1)} to {_g(st['x']['max'], 1)} {u}",
                        f"y range {_g(st['y']['min'], 1)} to {_g(st['y']['max'], 1)} {u}"], None, "60-bin histograms.")
    m = R["corr"]["matrix"]
    c["correlation"] = (["Pearson correlation of position and velocity components."],
                        [[""] + R["corr"]["labels"]] + [[R["corr"]["labels"][i]] + [_g(m[i][j], 2) for j in range(4)]
                                                        for i in range(4)], "Correlation matrix.")
    c["velocity"] = ([f"Peak speed {_g(R['speed']['peak'], 1)} {u}/s",
                      f"Max |ax| = {_g(max(abs(st['ax']['min']), abs(st['ax']['max'])), 1)} {u}/s2"],
                     None, "Central-difference derivatives of the filtered series.")
    c["energy"] = ([
        f"Energy proxy 0.5 v^2 + 0.5 w0^2 r^2 with w0 = 2 pi / T = {_g(E['omega0'], 3)} rad/s",
        f"Start {_g(E['E_start'], 1)} -> end {_g(E['E_end'], 1)} (arbitrary units)",
        f"Implied amplitude damping gamma = {_g(E['gamma_E'], 5)} 1/s (energy rate / 2)",
    ], None, "Not calibrated to physical energy.")
    c["residuals"] = ([
        f"Mean / SD residual x: {_g(RS['x']['mean'], 3)} / {_g(RS['x']['std'], 3)} {u}",
        f"Mean / SD residual y: {_g(RS['y']['mean'], 3)} / {_g(RS['y']['std'], 3)} {u}",
        f"Outliers replaced: {RS['outliers_x']} x, {RS['outliers_y']} y",
    ], None, "Residual = filtered - original series.")
    return c


def build_pptx(R):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    DARK, MUTED, LIGHT, BG = RGBColor(25, 35, 45), RGBColor(85, 100, 115), RGBColor(226, 234, 242), RGBColor(248, 250, 252)
    figs = render_all(R, dark=False)
    content = _slide_content(R)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def add_fit(slide, png, x, y, max_w, max_h):
        """Insert a PNG scaled to fit inside the box (never overflows the slide)."""
        from PIL import Image
        w, h = Image.open(io.BytesIO(png)).size
        k = min(max_w / (w / 130.0), max_h / (h / 130.0))
        slide.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), width=Inches(k * w / 130.0), height=Inches(k * h / 130.0))

    def base(title, sub=""):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        tb = s.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.55))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold, p.font.color.rgb = title, Pt(25), True, DARK
        p.alignment = PP_ALIGN.LEFT
        if sub:
            tb = s.shapes.add_textbox(Inches(0.58), Inches(0.82), Inches(12), Inches(0.3))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            p.text, p.font.size, p.font.color.rgb = sub, Pt(10.5), MUTED
            p.alignment = PP_ALIGN.LEFT
        return s

    def bullets(s, items, x, y, w, h, size=12):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text, p.font.size, p.font.color.rgb = "• " + item, Pt(size), DARK
            p.space_after = Pt(5)

    def table(s, rows, x, y, w, h, fs=9.5):
        tb = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tb.cell(r, c)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if r == 0 else RGBColor(255, 255, 255)
                for p in cell.text_frame.paragraphs:
                    p.font.size, p.font.color.rgb, p.font.bold = Pt(fs), DARK, r == 0

    # title slide
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    for text, y, size, bold, color in (("Pendulum Trajectory Analysis", 2.4, 36, True, DARK),
                                       (f"{R['name']} - {R['series_label']}", 3.45, 16, False, MUTED),
                                       (f"{R['filter_desc']} | {datetime.now():%Y-%m-%d %H:%M}", 4.0, 12, False, MUTED)):
        tb = s.shapes.add_textbox(Inches(0.6), Inches(y), Inches(12.1), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold, p.font.color.rgb = text, Pt(size), bold, color
        p.alignment = PP_ALIGN.CENTER

    for cid, title in CARD_TITLES.items():
        items, rows, note = content[cid]
        s = base(title, f"{R['series_label']} | {R['filter_desc']}")
        wide = cid in ("stats",)
        if wide:
            add_fit(s, figs[cid], 0.7, 1.2, 11.9, 4.9)
            bullets(s, items, 0.7, 6.3, 11.9, 0.8, 11)
        else:
            bullets(s, items, 0.55, 1.15, 5.3, 3.6 if rows else 5.6, 11.5)
            add_fit(s, figs[cid], 6.0, 1.15, 7.0, 5.9)
            if rows:
                table(s, rows, 0.6, 4.2 if len(rows) < 7 else 3.6, 5.2, 0.3 * len(rows), 9)
        try:
            s.notes_slide.notes_text_frame.text = note
        except Exception:  # noqa: BLE001
            pass

    s = base("Notes")
    bullets(s, [
        "Generated from the Flask dashboard; every number is recomputed from the current filter settings.",
        f"Series: {R['series_label']}; filter: {R['filter_desc']}; outliers: {R['params']['outlier']}.",
        "Velocity/acceleration are central differences - they amplify noise, so filter first.",
        "Energy values are proxies in arbitrary units, not calibrated physical energy.",
        "z(t) from apparent size is not true 3-D position.",
    ], 0.6, 1.3, 12, 4, 14)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
